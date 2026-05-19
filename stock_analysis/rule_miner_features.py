"""规则挖掘 Step 1: 特征工程 — 3,817笔往返交易 → 特征矩阵（parquet）

工作流：
1. 从 xlsx 加载原始交易 → FIFO 匹配 → 3,817 笔往返交易
2. 对每笔交易：读取买入日的 K 线数据计算特征
3. 打标签：盈利=1 / 亏损=0
4. 输出到 parquet 供规则挖掘用
"""

import json, os, sys
from collections import defaultdict
from datetime import datetime

import numpy as np
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import openpyxl

DATA_DIR = "D:/1989n/stock_data"
EASTMONEY_DIR = "D:/东方财富"
ML_DIR = os.path.join(DATA_DIR, "ml")
os.makedirs(ML_DIR, exist_ok=True)


# ============================================================
# Part 1: 加载原始交易数据（复用 bs_full_analysis.py 的逻辑）
# ============================================================

def load_new_format(filepath):
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[4] is None:
            continue
        trade_type = str(row[5])
        if trade_type not in ('证券买入', '证券卖出'):
            continue
        qty = int(row[6]) if row[6] else 0
        price = float(row[7]) if row[7] else 0
        amount = float(row[8]) if row[8] else 0
        net_amount = float(row[9]) if row[9] else 0
        commission = float(row[10]) if row[10] else 0
        stamp_tax = float(row[12]) if row[12] else 0
        transfer_fee = float(row[13]) if row[13] else 0
        if qty <= 0:
            continue
        rows.append({
            "date": str(row[1]) if row[1] else str(row[0]),
            "time": str(row[2]) if row[2] else "",
            "code": str(row[3]).strip().zfill(6),
            "name": str(row[4]),
            "direction": trade_type,
            "qty": qty,
            "price": price,
            "amount": amount,
            "net_amount": net_amount,
            "commission": commission,
            "stamp_tax": stamp_tax,
            "transfer_fee": transfer_fee,
        })
    wb.close()
    return rows


def load_old_format(filepath):
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[6] != "已成" or not row[8] or row[8] <= 0:
            continue
        direction = str(row[4])
        if "买入" not in direction and "卖出" not in direction:
            continue
        rows.append({
            "date": str(row[0]),
            "time": str(row[1]),
            "code": str(row[2]).strip().zfill(6),
            "name": str(row[3]),
            "direction": direction,
            "qty": int(row[5]),
            "price": float(row[10]) if row[10] else float(row[7]),
            "amount": float(row[9]),
            "net_amount": float(row[9]),
            "commission": 0,
            "stamp_tax": 0,
            "transfer_fee": 0,
        })
    wb.close()
    return rows


def load_all_trades():
    all_trades = []
    for f in ["2021-22.xlsx", "2022-2023.xlsx", "2023-2024xlsx.xlsx",
              "2024-2025.xlsx", "2025-2026.xlsx"]:
        path = os.path.join(EASTMONEY_DIR, f)
        if os.path.exists(path):
            trades = load_new_format(path)
            all_trades.extend(trades)
            print(f"  {f}: {len(trades)} 笔")
    old_path = os.path.join(EASTMONEY_DIR, "Table.xlsx")
    if os.path.exists(old_path):
        trades = load_old_format(old_path)
        all_trades.extend(trades)
        print(f"  Table.xlsx: {len(trades)} 笔")
    print(f"  总计: {len(all_trades)} 笔原始记录")
    return all_trades


# ============================================================
# Part 2: FIFO 匹配 → 往返交易
# ============================================================

def match_fifo(all_trades):
    by_stock = defaultdict(list)
    for t in all_trades:
        by_stock[t["code"]].append(t)
    for code in by_stock:
        by_stock[code].sort(key=lambda x: (x["date"], x["time"]))

    round_trips = []
    for code, trades in by_stock.items():
        buy_queue = []
        for t in trades:
            if "买入" in t["direction"]:
                buy_queue.append(t)
            elif "卖出" in t["direction"]:
                sell_qty = t["qty"]
                while sell_qty > 0 and buy_queue:
                    buy_t = buy_queue[0]
                    match_qty = min(sell_qty, buy_t["qty"])
                    buy_price = buy_t["price"]
                    sell_price = t["price"]

                    gross_profit = (sell_price - buy_price) * match_qty
                    buy_cost = (buy_t["commission"] + buy_t["stamp_tax"] + buy_t["transfer_fee"]) * (
                        match_qty / buy_t["qty"]
                    ) if buy_t["qty"] > 0 else 0
                    sell_cost = (t["commission"] + t["stamp_tax"] + t["transfer_fee"]) * (
                        match_qty / t["qty"]
                    ) if t["qty"] > 0 else 0
                    total_cost = buy_cost + sell_cost
                    net_profit = gross_profit - total_cost
                    gross_pct = ((sell_price - buy_price) / buy_price) * 100 if buy_price else 0

                    hold_days = 0
                    try:
                        bd = datetime.strptime(buy_t["date"][:10], "%Y-%m-%d")
                        sd = datetime.strptime(t["date"][:10], "%Y-%m-%d")
                        hold_days = (sd - bd).days
                    except ValueError:
                        pass

                    round_trips.append({
                        "code": code,
                        "name": buy_t["name"],
                        "buy_date": buy_t["date"][:10],
                        "sell_date": t["date"][:10],
                        "buy_price": round(buy_price, 2),
                        "sell_price": round(sell_price, 2),
                        "qty": match_qty,
                        "gross_profit": round(gross_profit, 2),
                        "net_profit": round(net_profit, 2),
                        "gross_pct": round(gross_pct, 2),
                        "hold_days": hold_days,
                        "buy_cost": round(buy_cost, 2),
                        "sell_cost": round(sell_cost, 2),
                    })

                    buy_t["qty"] -= match_qty
                    sell_qty -= match_qty
                    if buy_t["qty"] == 0:
                        buy_queue.pop(0)

    print(f"  FIFO匹配完成: {len(round_trips)} 笔往返交易")
    return round_trips


# ============================================================
# Part 3: K线特征计算
# ============================================================

def load_kline(code):
    """Load K-line parquet for a stock code."""
    path = os.path.join(DATA_DIR, "market_pool", "kline", f"{code}.parquet")
    if not os.path.exists(path):
        return None
    try:
        df = pq.read_table(path).to_pandas()
        if "date" in df.columns:
            df["date"] = df["date"].astype(str)
        df = df.sort_values("date").reset_index(drop=True)
        return df
    except Exception as e:
        print(f"    [warn] 读取 {code} K线失败: {e}")
        return None


def compute_ma(series, window):
    return series.rolling(window).mean()


def compute_rsi(series, period=14):
    delta = pd.Series(series).diff()
    gain = delta.clip(lower=0)
    loss = -delta.clip(upper=0)
    avg_gain = gain.rolling(window=period).mean()
    avg_loss = loss.rolling(window=period).mean()
    rs = avg_gain / avg_loss.replace(0, np.nan)
    rsi = 100 - (100 / (1 + rs))
    return rsi.values if hasattr(rsi, 'values') else rsi


def compute_macd(series, fast=12, slow=26, signal=9):
    series = pd.Series(series)
    ema_fast = series.ewm(span=fast).mean()
    ema_slow = series.ewm(span=slow).mean()
    dif = ema_fast - ema_slow
    dea = dif.ewm(span=signal).mean()
    macd_bar = 2 * (dif - dea)
    return dif.values if hasattr(dif, 'values') else dif, \
           dea.values if hasattr(dea, 'values') else dea, \
           macd_bar.values if hasattr(macd_bar, 'values') else macd_bar


def get_market_state(market_calendar, date_str):
    """Get market state for a given date."""
    if date_str in market_calendar:
        return market_calendar[date_str]
    return None


def compute_features(round_trips, market_calendar):
    """For each round-trip, compute features at buy point."""
    rows = []
    total = len(round_trips)
    cached_kline = {}

    for i, rt in enumerate(round_trips):
        if (i + 1) % 500 == 0:
            print(f"    特征计算: {i+1}/{total}")

        code = rt["code"]
        buy_date = rt["buy_date"]

        # Load K-line (cache)
        if code not in cached_kline:
            cached_kline[code] = load_kline(code)
        kline = cached_kline[code]
        if kline is None or len(kline) < 60:
            continue

        # Find buy date index in K-line
        # K-line dates are YYYYMMDD, buy_date is YYYY-MM-DD
        buy_date_kline = buy_date.replace("-", "")

        # Try exact match first
        if buy_date_kline in kline["date"].values:
            idx = kline[kline["date"] == buy_date_kline].index[0]
        else:
            # Try day before (if buy on non-trading day)
            from datetime import timedelta
            try:
                bd = datetime.strptime(buy_date, "%Y-%m-%d")
                for offset in range(1, 10):
                    test_date = (bd - timedelta(days=offset)).strftime("%Y%m%d")
                    if test_date in kline["date"].values:
                        idx = kline[kline["date"] == test_date].index[0]
                        break
                else:
                    continue
            except ValueError:
                continue

        if idx < 20:
            continue

        # === Compute features ===
        close = kline["close"].values
        close_series = kline["close"]

        # Price vs MA ratios
        ma5 = compute_ma(kline["close"], 5).values[idx]
        ma10 = compute_ma(kline["close"], 10).values[idx]
        ma20 = compute_ma(kline["close"], 20).values[idx]
        ma60 = compute_ma(kline["close"], 60).values[idx]
        price_at_buy = close[idx]

        # Volume ratio
        volume = kline["volume"].values
        avg_vol_20 = volume[max(0, idx - 20):idx].mean() if idx >= 20 else volume[:idx].mean()
        vol_at_buy = volume[idx]
        volume_ratio = vol_at_buy / avg_vol_20 if avg_vol_20 > 0 else 1.0

        # RSI
        rsi_values = compute_rsi(close_series)
        rsi_at_buy = rsi_values[idx]

        # MACD
        dif, dea, macd_bar = compute_macd(close_series)
        macd_state = "golden_cross" if dif[idx] > dea[idx] and dif[idx - 1] <= dea[idx - 1] else \
                     "dead_cross" if dif[idx] < dea[idx] and dif[idx - 1] >= dea[idx - 1] else \
                     "bullish" if dif[idx] > dea[idx] else "bearish"

        # Volatility (20-day)
        returns_20 = np.diff(close[max(0, idx - 20):idx + 1]) / close[max(0, idx - 20):idx]
        volatility_20 = float(np.std(returns_20) * 100) if len(returns_20) > 1 else 0

        # Price position in recent range
        high_20 = kline["high"].values[max(0, idx - 20):idx + 1].max()
        low_20 = kline["low"].values[max(0, idx - 20):idx + 1].min()
        price_position = (price_at_buy - low_20) / (high_20 - low_20) if (high_20 - low_20) > 0 else 0.5

        # Trend before buy (last 5 days return)
        pct_5d = ((close[idx] - close[max(0, idx - 5)]) / close[max(0, idx - 5)]) * 100 if idx >= 5 else 0

        # Market state at buy date
        mkt = get_market_state(market_calendar, buy_date)

        # Label
        is_win = 1 if rt["net_profit"] > 0 else 0

        # Normalize code to 6 digits
        code_normalized = code.zfill(6)

        row = {
            "code": code_normalized,
            "name": rt["name"],
            "buy_date": buy_date,
            "sell_date": rt["sell_date"],
            "hold_days": rt["hold_days"],
            "gross_pct": rt["gross_pct"],
            "net_profit": rt["net_profit"],
            "label": is_win,

            # Price vs MA
            "price_ma5_ratio": round(price_at_buy / ma5, 4) if ma5 and ma5 > 0 else 1.0,
            "price_ma10_ratio": round(price_at_buy / ma10, 4) if ma10 and ma10 > 0 else 1.0,
            "price_ma20_ratio": round(price_at_buy / ma20, 4) if ma20 and ma20 > 0 else 1.0,
            "price_ma60_ratio": round(price_at_buy / ma60, 4) if ma60 and ma60 > 0 else 1.0,

            # Volume
            "volume_ratio": round(volume_ratio, 4),

            # RSI
            "rsi_14": round(rsi_at_buy, 2) if not np.isnan(rsi_at_buy) else 50,

            # MACD
            "macd_state": macd_state,

            # Volatility & position
            "volatility_20": round(volatility_20, 4),
            "price_position_20": round(price_position, 4),
            "pct_5d_before_buy": round(pct_5d, 2),

            # Market context
            "market_trend": mkt.get("trend", "unknown") if mkt else "unknown",
            "market_volume_state": mkt.get("volume_state", "unknown") if mkt else "unknown",
            "market_sentiment": mkt.get("sentiment", "unknown") if mkt else "unknown",
            "market_pct_chg": mkt.get("pct_chg", 0) if mkt else 0,
        }
        rows.append(row)

    return pd.DataFrame(rows)


# ============================================================
# Main
# ============================================================

def main():
    print("=" * 50)
    print("规则挖掘 Step 1: 特征工程")
    print("=" * 50)

    # 1. Load trades
    print("\n[1/4] 加载原始交易数据...")
    all_trades = load_all_trades()

    # 2. FIFO match
    print("\n[2/4] FIFO 匹配往返交易...")
    round_trips = match_fifo(all_trades)
    print(f"  胜率: {sum(1 for r in round_trips if r['net_profit'] > 0) / len(round_trips) * 100:.1f}%")
    print(f"  总净利: {sum(r['net_profit'] for r in round_trips):,.0f}")

    # Save raw round trips for inspection
    rt_path = os.path.join(ML_DIR, "round_trips.json")
    with open(rt_path, "w", encoding="utf-8") as f:
        json.dump(round_trips[:100], f, ensure_ascii=False, indent=2)
    print(f"  前100笔已保存到: {rt_path}")

    # 3. Load market calendar
    print("\n[3/4] 加载市场日历...")
    mcal_path = os.path.join(DATA_DIR, "market_calendar_full.json")
    with open(mcal_path) as f:
        market_calendar = json.load(f)
    print(f"  加载了 {len(market_calendar)} 个交易日")

    # 4. Compute features
    print("\n[4/4] 计算K线特征...")
    df = compute_features(round_trips, market_calendar)
    print(f"  特征矩阵: {len(df)} 行, {len(df.columns)} 列")
    print(f"  label分布: 盈利={df['label'].sum()}, 亏损={len(df)-df['label'].sum()}")

    # Save to parquet
    out_path = os.path.join(ML_DIR, "trade_features.parquet")
    df.to_parquet(out_path, index=False)
    print(f"\n✅ 特征矩阵已保存: {out_path} ({os.path.getsize(out_path) / 1024:.0f} KB)")

    # Summary stats
    print("\n简单统计:")
    print(f"  平均持仓天数: {df['hold_days'].mean():.1f}")
    print(f"  平均 gross_pct: {df['gross_pct'].mean():.2f}%")
    print(f"  平均 volume_ratio: {df['volume_ratio'].mean():.2f}")
    print(f"  平均 rsi_14: {df['rsi_14'].mean():.1f}")
    print(f"  平均 price_ma20_ratio: {df['price_ma20_ratio'].mean():.4f}")


if __name__ == "__main__":
    main()
