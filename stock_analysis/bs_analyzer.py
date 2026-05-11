"""
BS点交易分析 — 匹配买卖对、计算盈亏、分类分析
"""
import json
import os
from collections import defaultdict
from datetime import datetime
import openpyxl

DATA_DIR = "D:/1989n/stock_data"
BS_FILE = "D:/1989n/Table.xlsx"


def load_bs_data(filepath):
    """加载东方财富BS点数据"""
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    headers = [c.value for c in list(ws.iter_rows(min_row=1, max_row=1))[0]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[6] == "已成" and row[8] and row[8] > 0:  # 只取已成且有成交数量
            rows.append(dict(zip(headers, row)))
    print(f"加载 {len(rows)} 条已成交易记录")
    return rows


def match_trades_fifo(trades):
    """FIFO匹配买卖对，返回完整往返交易"""
    # 按股票分组
    by_stock = defaultdict(list)
    for t in trades:
        code = str(t["证券代码"])
        by_stock[code].append(t)

    round_trips = []
    unmatched_buys = []

    for code, stock_trades in by_stock.items():
        # 按日期时间排序
        stock_trades.sort(key=lambda x: (str(x["委托日期"]), str(x["委托时间"])))

        buy_queue = []  # (日期, 数量, 价格, 金额, 名称)
        for t in stock_trades:
            direction = t["委托方向"]
            qty = int(t["成交数量"])
            price = float(t["成交均价"])
            amount = float(t["成交金额"])
            date = str(t["委托日期"])
            name = str(t["证券名称"])
            market = str(t["交易市场"])

            if "买入" in direction:
                buy_queue.append({
                    "date": date, "qty": qty, "price": price,
                    "amount": amount, "name": name, "code": code, "market": market
                })
            elif "卖出" in direction:
                sell_qty = qty
                while sell_qty > 0 and buy_queue:
                    buy = buy_queue[0]
                    if buy["qty"] <= sell_qty:
                        # 完全消耗这笔买入
                        matched_qty = buy["qty"]
                        sell_qty -= matched_qty
                        buy_queue.pop(0)
                    else:
                        # 部分消耗
                        matched_qty = sell_qty
                        buy["qty"] -= sell_qty
                        buy["amount"] = buy["amount"] * (buy["qty"] / (buy["qty"] + sell_qty))
                        sell_qty = 0

                    buy_amount = matched_qty * buy["price"]
                    sell_amount = matched_qty * price
                    gross_profit = sell_amount - buy_amount
                    gross_pct = (price - buy["price"]) / buy["price"] * 100
                    hold_days = (datetime.strptime(date, "%Y-%m-%d") -
                                 datetime.strptime(buy["date"], "%Y-%m-%d")).days

                    round_trips.append({
                        "code": code,
                        "name": name,
                        "market": market,
                        "buy_date": buy["date"],
                        "sell_date": date,
                        "buy_price": round(buy["price"], 3),
                        "sell_price": round(price, 3),
                        "qty": matched_qty,
                        "buy_amount": round(buy_amount, 2),
                        "sell_amount": round(sell_amount, 2),
                        "gross_profit": round(gross_profit, 2),
                        "gross_pct": round(gross_pct, 2),
                        "hold_days": hold_days,
                    })

    # 收集未匹配的买入（还在持仓）
    for code, stock_trades in by_stock.items():
        stock_trades.sort(key=lambda x: (str(x["委托日期"]), str(x["委托时间"])))
        buy_queue = []
        for t in stock_trades:
            if "买入" in str(t["委托方向"]):
                buy_queue.append({
                    "date": str(t["委托日期"]), "qty": int(t["成交数量"]),
                    "price": float(t["成交均价"]), "name": str(t["证券名称"]),
                    "code": code
                })
            elif "卖出" in str(t["委托方向"]):
                sell_qty = int(t["成交数量"])
                while sell_qty > 0 and buy_queue:
                    buy = buy_queue[0]
                    matched = min(buy["qty"], sell_qty)
                    buy["qty"] -= matched
                    sell_qty -= matched
                    if buy["qty"] <= 0:
                        buy_queue.pop(0)

        for b in buy_queue:
            if b["qty"] > 0:
                unmatched_buys.append(b)

    return round_trips, unmatched_buys


def classify_trades(round_trips):
    """分类：赚/亏/平"""
    winners = [t for t in round_trips if t["gross_pct"] > 0.5]
    losers = [t for t in round_trips if t["gross_pct"] < -0.5]
    flats = [t for t in round_trips if -0.5 <= t["gross_pct"] <= 0.5]

    return winners, losers, flats


def analyze_category(trades, label):
    """分析某一类交易的特征"""
    if not trades:
        return {"label": label, "count": 0}

    total_profit = sum(t["gross_profit"] for t in trades)
    avg_pct = sum(t["gross_pct"] for t in trades) / len(trades)
    avg_hold = sum(t["hold_days"] for t in trades) / len(trades)
    median_pct = sorted([t["gross_pct"] for t in trades])[len(trades) // 2]

    # 按月份分布
    by_month = defaultdict(lambda: {"count": 0, "profit": 0})
    for t in trades:
        month = t["buy_date"][:7]
        by_month[month]["count"] += 1
        by_month[month]["profit"] += t["gross_profit"]

    # 高频股票
    stock_freq = defaultdict(lambda: {"count": 0, "profit": 0, "name": ""})
    for t in trades:
        code = t["code"]
        stock_freq[code]["count"] += 1
        stock_freq[code]["profit"] += t["gross_profit"]
        stock_freq[code]["name"] = t["name"]

    top_worst = sorted(stock_freq.items(), key=lambda x: x[1]["profit"])[:10]
    top_best = sorted(stock_freq.items(), key=lambda x: x[1]["profit"], reverse=True)[:10]

    return {
        "label": label,
        "count": len(trades),
        "total_profit": round(total_profit, 2),
        "avg_pct": round(avg_pct, 2),
        "median_pct": round(median_pct, 2),
        "avg_hold_days": round(avg_hold, 1),
        "monthly": dict(by_month),
        "top_best_stocks": [(code, data) for code, data in top_best],
        "top_worst_stocks": [(code, data) for code, data in top_worst],
    }


def main():
    trades = load_bs_data(BS_FILE)
    round_trips, unmatched = match_trades_fifo(trades)
    winners, losers, flats = classify_trades(round_trips)

    print(f"\n=== BS点分析结果 ===")
    print(f"总成交记录: {len(trades)}")
    print(f"完整往返交易: {len(round_trips)}")
    print(f"盈利交易: {len(winners)} ({len(winners)/len(round_trips)*100:.1f}%)")
    print(f"亏损交易: {len(losers)} ({len(losers)/len(round_trips)*100:.1f}%)")
    print(f"持平交易: {len(flats)} ({len(flats)/len(round_trips)*100:.1f}%)")
    print(f"未平仓买入: {len(unmatched)}")

    total_pnl = sum(t["gross_profit"] for t in round_trips)
    print(f"\n总盈亏: {total_pnl:,.2f}")
    print(f"总盈利金额: {sum(t['gross_profit'] for t in winners):,.2f}")
    print(f"总亏损金额: {sum(t['gross_profit'] for t in losers):,.2f}")

    # 详细分析
    results = {
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "summary": {
            "total_trades": len(trades),
            "round_trips": len(round_trips),
            "winners": len(winners),
            "losers": len(losers),
            "flats": len(flats),
            "total_pnl": round(total_pnl, 2),
            "win_rate": round(len(winners) / len(round_trips) * 100, 1) if round_trips else 0,
        },
        "winners": analyze_category(winners, "盈利"),
        "losers": analyze_category(losers, "亏损"),
        "flats": analyze_category(flats, "持平"),
        "unmatched_buys": unmatched,
    }

    # 保存
    out_path = os.path.join(DATA_DIR, "bs_analysis.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2, default=str)
    print(f"\n结果已保存到 {out_path}")

    # 打印每类详情
    for cat in [results["winners"], results["losers"], results["flats"]]:
        if cat["count"] == 0:
            continue
        print(f"\n--- {cat['label']} ({cat['count']}笔) ---")
        print(f"  总盈亏: {cat['total_profit']:,.2f}")
        print(f"  平均收益率: {cat['avg_pct']:.2f}%")
        print(f"  中位收益率: {cat['median_pct']:.2f}%")
        print(f"  平均持仓: {cat['avg_hold_days']:.1f}天")
        print(f"  最佳5只:")
        for code, data in cat["top_best_stocks"][:5]:
            print(f"    {code} {data['name']}: {data['count']}次, {data['profit']:,.2f}")
        print(f"  最差5只:")
        for code, data in cat["top_worst_stocks"][:5]:
            print(f"    {code} {data['name']}: {data['count']}次, {data['profit']:,.2f}")

    return results


if __name__ == "__main__":
    main()
